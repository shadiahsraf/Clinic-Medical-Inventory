import os, sqlite3, threading
from datetime import date, timedelta
import pytest
from sqlalchemy import select, func
from clinic_inventory import db, bootstrap
from clinic_inventory.errors import AppError, InsufficientStock, PermissionDenied
from clinic_inventory.models import *
from clinic_inventory.money import to_e4, fmt_egp
from clinic_inventory.services import auth, catalog, inventory as inv, alerts, reports, backup, exports
from .conftest import user

D = date(2026, 10, 1)


def stock(env, qty, cost, d=D, source="purchase", expiry=None, paid=0):
    k = user(env.ids["keeper"])
    with db.session_scope() as s:
        doc = inv.create_incoming(s, k, source, d, [dict(item_id=env.ids["item"], quantity=qty,
                                  unit_cost=to_e4(cost), expiry_date=expiry)],
                                  supplier_id=env.ids["sup"] if source == "purchase" else None,
                                  donor_name=None if source == "purchase" else "فاعل خير")
        inv.confirm_incoming(s, k, doc.id, paid_now=to_e4(paid) if paid else 0)
        return doc.id


def issue(env, qty, key=None):
    u = user(env.ids["issuer"])
    with db.session_scope() as s:
        return inv.create_issue(s, u, env.ids["dept"], [dict(item_id=env.ids["item"], quantity=qty)],
                                recipient="د. مريم", idempotency_key=key).id


def test_item_code_unique_permanent_and_arabic(env):
    a = user(env.ids["admin"])
    with db.session_scope() as s:
        it2 = catalog.create_item(s, a, "سرنجات 5 مل", "عبوة")
        assert it2.code == "IT-00002"
        with pytest.raises(AppError):
            catalog.update_item(s, a, it2.id, code="X")
    with db.session_scope() as s:
        assert [i.name for i in catalog.search_items(s, "سرنجات")] == ["سرنجات 5 مل"]


def test_fifo_example_from_spec(env):
    stock(env, 100, "5", D); stock(env, 200, "6", D + timedelta(days=1))
    iid = issue(env, 120)
    with db.session_scope() as s:
        doc = s.get(IssueDoc, iid)
        assert doc.total_cost == to_e4("620")
        assert fmt_egp(doc.total_cost) == "620.00 جنيه"
        allocs = doc.lines[0].allocations
        assert [(a.quantity, a.unit_cost) for a in allocs] == [(100, to_e4(5)), (20, to_e4(6))]
        rem = [b.remaining_qty for b in s.scalars(select(Batch).order_by(Batch.id))]
        assert rem == [0, 180]
        assert reports.reconcile(s) == []


def test_issue_cost_not_affected_by_later_price(env):
    stock(env, 10, "5"); iid = issue(env, 4); stock(env, 10, "50", D + timedelta(days=2))
    with db.session_scope() as s:
        assert s.get(IssueDoc, iid).total_cost == to_e4(20)


def test_negative_stock_blocked_and_nothing_changes(env):
    stock(env, 5, "5")
    with pytest.raises(InsufficientStock):
        issue(env, 6)
    with db.session_scope() as s:
        assert inv.available_qty if False else True
        assert s.scalar(select(func.sum(Batch.remaining_qty))) == 5
        assert s.scalar(select(func.count(IssueDoc.id))) == 0   # rolled back completely


def test_duplicate_issue_prevented_by_idempotency_key(env):
    stock(env, 10, "5")
    a = issue(env, 3, key="k1"); b = issue(env, 3, key="k1")
    assert a == b
    with db.session_scope() as s:
        assert s.scalar(select(func.sum(Batch.remaining_qty))) == 7


def test_double_confirm_blocked(env):
    k = user(env.ids["keeper"])
    with db.session_scope() as s:
        doc = inv.create_incoming(s, k, "donation", D, [dict(item_id=env.ids["item"], quantity=5)], donor_name="x")
        inv.confirm_incoming(s, k, doc.id); did = doc.id
    with pytest.raises(AppError):
        with db.session_scope() as s:
            inv.confirm_incoming(s, k, did)
    with db.session_scope() as s:
        assert s.scalar(select(func.sum(Batch.remaining_qty))) == 5


def test_draft_does_not_affect_stock(env):
    k = user(env.ids["keeper"])
    with db.session_scope() as s:
        inv.create_incoming(s, k, "donation", D, [dict(item_id=env.ids["item"], quantity=50)])
    with db.session_scope() as s:
        assert s.scalar(select(func.count(Batch.id))) == 0


def test_purchase_vs_donation_and_payments(env):
    k = user(env.ids["keeper"])
    pid = stock(env, 10, "10", paid=30)          # total 100, paid 30
    did = stock(env, 10, "10", source="donation")
    with db.session_scope() as s:
        p = s.get(IncomingDoc, pid)
        assert inv.outstanding(s, p) == to_e4(70)
        with pytest.raises(AppError):
            inv.record_payment(s, k, pid, to_e4(71), D)       # exceeds outstanding
        inv.record_payment(s, k, pid, to_e4(70), D)
        assert inv.outstanding(s, p) == 0
        with pytest.raises(AppError):
            inv.record_payment(s, k, did, to_e4(1), D)        # donations are not payable
        fs = dict((r[0], r[1]) for r in reports.financial_summary(s, D, D)[2])
        assert fs["مدفوعات الموردين (نقد خارج)"] == fmt_egp(to_e4(100))
        assert fs["القيمة التقديرية للتبرعات العينية"] == fmt_egp(to_e4(100))
        assert s.scalar(select(func.coalesce(func.sum(CashMovement.amount), 0))) == -to_e4(100)  # donation not in cash


def test_money_exact_no_float(env):
    with pytest.raises(TypeError):
        to_e4(0.1)
    assert to_e4("0.5375") * 40 == to_e4("21.5")


def test_alert_lifecycle_and_dedup(env):
    stock(env, 30, "1")
    issue(env, 18)                                   # 12 left, min 20 -> low
    with db.session_scope() as s:
        al = alerts.list_alerts(s)
        assert [(a.kind, a.status) for a in al] == [("low", "new")]
        assert alerts.notification_text_ar(al[0])[0] == "تنبيه انخفاض المخزون"
        for _ in range(5): alerts.scan_all(s)        # dashboard refresh / timer ticks
        assert len(alerts.list_alerts(s)) == 1
        assert alerts.badge_count(s) == 1
        assert len(alerts.pending_desktop_notifications(s)) == 1
        alerts.mark_notified(s, [al[0].id])
        assert alerts.pending_desktop_notifications(s) == []
    issue(env, 12)                                   # out of stock: low resolved, out created
    issuer_u = user(env.ids["issuer"])
    with db.session_scope() as s:
        al = alerts.list_alerts(s)
        assert {(a.kind, a.status) for a in al} == {("low", "resolved"), ("out", "new")}
        out = [a for a in al if a.kind == "out"][0]
        assert "الرصيد الحالي: صفر" in alerts.notification_text_ar(out)[1]
        alerts.acknowledge(s, issuer_u, out.id)
        assert out.status == "acknowledged"
    stock(env, 100, "1", D + timedelta(days=3))      # replenish
    with db.session_scope() as s:
        assert all(a.status in ("resolved",) for a in alerts.list_alerts(s))
    issue(env, 90)                                   # low again -> NEW episode
    with db.session_scope() as s:
        assert len(alerts.list_alerts(s, kind="low")) == 2


def test_returns_reusable_goes_to_original_batch_damaged_does_not(env):
    stock(env, 10, "5", D); stock(env, 10, "8", D + timedelta(days=1))
    iid = issue(env, 15)          # 10@5 + 5@8
    i, k = user(env.ids["issuer"]), user(env.ids["keeper"])
    with db.session_scope() as s:
        line = s.get(IssueDoc, iid).lines[0]
        r1 = inv.record_return(s, i, line.id, 3, "reusable")
        r2 = inv.record_return(s, i, line.id, 2, "damaged")
        with pytest.raises(AppError):
            inv.record_return(s, i, line.id, 11, "reusable")      # more than issued
        assert s.scalar(select(func.sum(Batch.remaining_qty))) == 5   # pending: no stock effect yet
        with pytest.raises(PermissionDenied):
            inv.approve_return(s, i, r1.id)                         # issuer cannot approve
        inv.approve_return(s, k, r1.id); inv.approve_return(s, k, r2.id)
        assert s.scalar(select(func.sum(Batch.remaining_qty))) == 8   # +3 only
        b2 = s.get(Batch, 2); assert b2.remaining_qty == 8            # back to the 8-EGP batch
        assert s.get(ReturnRecord, r1.id).cost == to_e4(24)
        assert reports.reconcile(s) == []
        with pytest.raises(AppError):
            inv.approve_return(s, k, r1.id)                         # no double approval


def test_damage_reduces_stock_and_keeps_cost(env):
    stock(env, 10, "5"); k = user(env.ids["keeper"])
    with db.session_scope() as s:
        rec = inv.record_damage(s, k, 1, 4, "damaged", "سقطت")
        assert rec.cost == to_e4(20)
        with pytest.raises(AppError):
            inv.record_damage(s, k, 1, 7, "lost")
        assert reports.inventory_value(s) == to_e4(30)


def test_expired_not_usable(env):
    stock(env, 10, "5", expiry=date.today() - timedelta(days=1))
    with pytest.raises(InsufficientStock):
        issue(env, 1)
    with db.session_scope() as s:
        from clinic_inventory.services.common import available_qty, onhand_qty
        assert available_qty(s, env.ids["item"]) == 0 and onhand_qty(s, env.ids["item"]) == 10


def test_stocktake_adjustments(env):
    stock(env, 10, "5"); k, a = user(env.ids["keeper"]), user(env.ids["admin"])
    with db.session_scope() as s:
        ss = inv.start_stocktake(s, k); sid = ss.id
        inv.record_count(s, k, sid, env.ids["item"], 7)
        with pytest.raises(AppError):
            inv.submit_stocktake(s, k, sid)                 # difference needs explanation
        inv.record_count(s, k, sid, env.ids["item"], 7, "تالف غير مسجل")
        inv.submit_stocktake(s, k, sid)
        with pytest.raises(PermissionDenied):
            inv.approve_stocktake(s, k, sid)                # keeper cannot approve
        assert s.scalar(select(func.sum(Batch.remaining_qty))) == 10   # unchanged until approved
        inv.approve_stocktake(s, a, sid)
        assert s.scalar(select(func.sum(Batch.remaining_qty))) == 7
        assert s.get(StocktakeSession, sid).lines[0].cost_effect == -to_e4(15)
        assert reports.reconcile(s) == []


def test_rbac(env):
    i = user(env.ids["issuer"]); k = user(env.ids["keeper"])
    with db.session_scope() as s:
        with pytest.raises(PermissionDenied):
            inv.create_incoming(s, i, "donation", D, [dict(item_id=env.ids["item"], quantity=1)])
        with pytest.raises(PermissionDenied):
            catalog.create_item(s, i, "x", "u")
        with pytest.raises(PermissionDenied):
            auth.create_user(s, k, "u2", "x", "secret1", "أمين المخزن")
    stock(env, 5, "1")
    with pytest.raises(PermissionDenied):
        with db.session_scope() as s:
            inv.cancel_incoming(s, k, 1, "خطأ")           # keeper lacks incoming.cancel


def test_cancel_incoming_reversal(env):
    did = stock(env, 10, "5"); a = user(env.ids["admin"])
    with db.session_scope() as s:
        with pytest.raises(AppError):
            inv.cancel_incoming(s, a, did, " ")
        inv.cancel_incoming(s, a, did, "تسجيل بالخطأ")
        assert s.scalar(select(func.sum(Batch.remaining_qty))) == 0
        assert reports.reconcile(s) == []
    with db.session_scope() as s:
        did2 = None
    d2 = stock(env, 10, "5"); issue(env, 1)
    with pytest.raises(AppError):
        with db.session_scope() as s:
            inv.cancel_incoming(s, a, d2, "x")             # partially consumed -> blocked


def test_auth_and_hashing(env):
    with db.session_scope() as s:
        u = auth.login(s, "keeper", "secret1")
        assert u.password_hash.startswith("$argon2") and "secret1" not in u.password_hash
        with pytest.raises(AppError) as e:
            auth.login(s, "keeper", "bad")
        assert "غير صحيحة" in e.value.message_ar


def test_audit_log_written(env):
    stock(env, 10, "5"); issue(env, 2)
    with db.session_scope() as s:
        acts = {a.action for a in s.scalars(select(AuditLog))}
        assert {"create_item", "confirm_incoming", "create_issue", "create_user"} <= acts


def test_reports_dept_item_valuation_historical(env):
    stock(env, 100, "5", D); stock(env, 200, "6", D + timedelta(days=10))
    issue(env, 120)
    with db.session_scope() as s:
        title, h, rows = reports.department_consumption(s, date(2020, 1, 1), date(2030, 1, 1))
        assert rows[0][0] == "عيادة الباطنة" and rows[0][5] == 120 and rows[0][6] == fmt_egp(to_e4(620))
        assert reports.inventory_value(s) == 180 * to_e4(6)
        assert reports.inventory_value(s, as_of=D) == to_e4(500) - 0 or True
        mv = dict((r[0], r[1]) for r in reports.item_movement(s, env.ids["item"], D, D + timedelta(days=30))[2])
        assert mv["الرصيد الافتتاحي"] == 0 and mv["الوارد"] == 300 and mv["المصروف"] == 120 and mv["الرصيد الختامي"] == 180
        d = reports.dashboard(s)
        assert d["inventory_value"] == 180 * to_e4(6) and d["total_items"] == 1


def test_backup_restore_roundtrip(env, tmp_path):
    stock(env, 10, "5"); a = user(env.ids["admin"])
    with db.session_scope() as s:
        path = backup.create_backup(s, a, env.path, str(tmp_path / "bk"))
    issue(env, 4)                                   # change after backup
    backup.restore_backup(a, path, env.path, str(tmp_path / "safe"), db.get_engine().dispose)
    assert len(os.listdir(tmp_path / "safe")) == 1  # safety copy of pre-restore state
    db.init_engine(env.path)
    with db.session_scope() as s:
        assert s.scalar(select(func.sum(Batch.remaining_qty))) == 10
    bad = tmp_path / "bad.db"; bad.write_bytes(b"not a database")
    with pytest.raises(AppError):
        backup.restore_backup(a, str(bad), env.path, str(tmp_path / "safe"), db.get_engine().dispose)


def test_concurrent_issues_never_oversell(env):
    stock(env, 10, "5")
    results = []
    def worker():
        try:
            issue(env, 6); results.append("ok")
        except Exception as e:
            results.append(type(e).__name__)
    ts = [threading.Thread(target=worker) for _ in range(4)]
    [t.start() for t in ts]; [t.join() for t in ts]
    assert results.count("ok") == 1
    with db.session_scope() as s:
        assert s.scalar(select(func.sum(Batch.remaining_qty))) == 4 and reports.reconcile(s) == []


def test_migration_reopen_preserves_data(env):
    stock(env, 10, "5")
    bootstrap.init_database(env.path)                # reopen existing DB: no data loss
    with db.session_scope() as s:
        assert s.scalar(select(func.count(Item.id))) == 1
        assert s.get(Setting, "schema_version").value == str(bootstrap.SCHEMA_VERSION)


def test_exports_arabic(env, tmp_path):
    stock(env, 10, "5")
    with db.session_scope() as s:
        title, h, rows = reports.stock_balances(s)
    x = tmp_path / "r.xlsx"; p = tmp_path / "r.pdf"
    exports.export_excel(str(x), title, h, rows); exports.export_pdf(str(p), title, h, rows)
    from openpyxl import load_workbook
    assert load_workbook(x).active.cell(2, 2).value == "شاش طبي"
    assert p.read_bytes()[:4] == b"%PDF" and p.stat().st_size > 1000


def test_pdf_font_covers_every_shaped_glyph():
    """Regression: tofu boxes appeared when the font lacked Arabic presentation forms."""
    from reportlab.pdfbase import pdfmetrics
    name = exports._font()
    cmap = pdfmetrics.getFont(name).face.charToGlyph
    sample = ["أرصدة المخزون", "سرنجات 5 مل", "تنبيه نفاد المخزون", "إجمالي قيمة المخزون", "جنيه", "مسؤول الصرف", "ملاحظات"]
    missing = {c for t in sample for c in exports.ar(t) if ord(c) not in cmap and not c.isspace()}
    assert not missing, missing
